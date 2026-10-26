
import streamlit as st
import fitz
import pytesseract
from docx import Document
from PIL import Image
import io
import re
import html

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

st.set_page_config(
    page_title="Medical Report Simplifier",
    page_icon="app_icon.png",
    layout="wide"
)




st.title("🩺 Medical Report Simplifier")

st.subheader(
    "Understand your medical reports in simple language"
)

st.markdown(
    """
    Welcome! Upload your medical report to understand
    medical terms and results in simple language.
    """
)

st.info(
    "For educational use only. This tool does not provide "
    "a diagnosis or replace advice from a qualified "
    "healthcare professional."
)


st.markdown(
    """
    <style>
    /* Main website background */
    .stApp {
        background: linear-gradient(
            135deg, #D9F0F7 0%, #EEF8FC 55%, #D5EAF3 100%
        );
    }

    /* Main content area */
    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }

    /* Main heading */
    h1 {
        color: #07566B !important;
        font-weight: 800 !important;
    }

    /* Other headings */
    h2, h3 {
        color: #087E8B !important;
    }

    /* Normal text */
    .stMarkdown p {
        color: #244653;
    }

    
    /* Success message - matching blue theme */
    div[data-testid="stAlert"][data-st-alert-type="success"] {
        background-color: #D9F0F7 !important;
        border: 1px solid #70B6CA !important;
        border-radius: 10px !important;
    }

    div[data-testid="stAlert"][data-st-alert-type="success"] p {
        color: #07566B !important;
        -webkit-text-fill-color: #07566B !important;
        font-weight: 600 !important;
    }

    /* File upload area */
    div[data-testid="stFileUploader"] {
        background-color: #FFFFFF;
        border: 2px dashed #168A9B;
        border-radius: 16px;
        padding: 20px;
    }

   
    /* Download button */
    .stDownloadButton > button {
        background-color: #087E8B !important;
        color: white !important;
        border-radius: 10px !important;
        border: none !important;
        padding: 0.5rem 1.2rem !important;
    }

    .stDownloadButton > button:hover {
        background-color: #07566B !important;
        color: white !important;
        border: 1px solid #07566B !important;
    }


    /* Original medical report text box */
    div[data-testid="stTextArea"] textarea {
        background-color: #E8F5FA !important;
        color: #164E63 !important;
        border: 2px solid #70B6CA !important;
        border-radius: 10px !important;
        font-size: 16px !important;
        line-height: 1.7 !important;
    }

   
    /* Fix original report text visibility */
    div[data-testid="stTextArea"] textarea {
        background-color: #E8F5FA !important;
        color: #164E63 !important;
        -webkit-text-fill-color: #164E63 !important;
        border: 2px solid #70B6CA !important;
        border-radius: 10px !important;
        font-size: 16px !important;
    }


   
    /* Expander background and heading */
    div[data-testid="stExpander"] {
        background-color: #EAF6FB !important;
        border: 1px solid #70B6CA !important;
        border-radius: 12px !important;
    }

    div[data-testid="stExpander"] summary,
    div[data-testid="stExpander"] summary p,
    div[data-testid="stExpander"] summary span {
        background-color: #EAF6FB !important;
        color: #07566B !important;
        -webkit-text-fill-color: #07566B !important;
        font-weight: bold !important;
    }


    </style>
    """,
    unsafe_allow_html=True
)

st.write(
    "Upload a PDF, Word document, or image to get easy-to-understand "
    "explanations of medical terms."
)
st.info(
    "For educational use only. This tool does not provide a diagnosis "
    "or replace advice from a qualified healthcare professional."
)


def simplify_report(text):
    
    explanations = {
        "hemoglobin":
            "Hemoglobin is a protein in red blood cells that carries oxygen.",
        "platelet count":
            "Platelets help blood clot and stop bleeding.",
        "white blood cell count":
            "White blood cells help the body fight infections.",
        "reference range":
            "A reference range shows values commonly expected by the laboratory.",
        "red blood cell count":
            "Red blood cells carry oxygen from the lungs to the body.",
        "hematocrit":
            "Hematocrit is the percentage of blood made up of red blood cells.",
        "white blood cells":
            "White blood cells help protect the body against infections.",
        "blood glucose":
            "Blood glucose is the amount of sugar in your blood.",
        "fasting glucose":
            "Fasting glucose measures blood sugar after not eating for a period.",
        "cholesterol":
            "Cholesterol is a fatty substance in the blood that the body needs in limited amounts.",
        "triglycerides":
            "Triglycerides are a type of fat found in the blood.",
        "creatinine":
            "Creatinine is a waste product that the kidneys help remove from the blood.",
        "urea":
            "Urea is a waste product removed from the blood by the kidneys.",
        "sodium":
            "Sodium is a mineral that helps control fluid balance and nerve function.",
        "potassium":
            "Potassium is a mineral important for muscles, nerves, and heart function.",
        "calcium":
            "Calcium helps keep bones and teeth strong and supports muscle function.",
        "vitamin d":
            "Vitamin D helps the body absorb calcium and supports bone health.",
        "vitamin b12":
            "Vitamin B12 helps the body make red blood cells and maintain nerve function.",
        "tsh":
            "TSH is a hormone test used to help assess thyroid function.",
        "heart rate":
            "Heart rate is the number of times your heart beats per minute.",
        "blood pressure":
            "Blood pressure measures the force of blood against artery walls.",
        "oxygen saturation":
            "Oxygen saturation estimates how much oxygen the blood is carrying.",
        "hba1c":
            "HbA1c estimates average blood sugar over roughly the past 2–3 months.",
        "bilirubin":
            "Bilirubin is a yellow substance produced when red blood cells break down.",
        "albumin":
            "Albumin is a protein in the blood that helps maintain fluid balance.",
        "alt":
            "ALT is an enzyme often checked when assessing liver health.",
        "ast":
            "AST is an enzyme that may be measured when assessing liver or other tissue health.",
        "uric acid":
            "Uric acid is a waste product formed when the body breaks down purines.",
        "urine protein":
            "Urine protein testing checks whether protein is present in the urine."
    }


    results = []

    for term, explanation in explanations.items():
        if term in text.lower():
            results.append(f"**{term.title()}:** {explanation}")

    if not results:
        results.append("No known medical terms were found.")

    return results


def check_numbers(original, simplified):
    original_numbers = re.findall(
        r"\b\d+(?:\.\d+)?\b", original
    )
    simplified_numbers = re.findall(
        r"\b\d+(?:\.\d+)?\b", simplified
    )

    return [
        number for number in simplified_numbers
        if number not in original_numbers
    ]





st.markdown(
    """
    <div style="
        background-color: #FFFFFF;
        border: 2px solid #168A9B;
        border-radius: 16px;
        padding: 22px;
        margin-top: 20px;
        margin-bottom: 10px;
    ">
        <h3 style="color: #07566B; margin-top: 0;">
            📄 Upload Your Medical Report
        </h3>
        <p style="color: #42616B;">
            Choose a report from your computer to begin.
        </p>
        <p style="color: #087E8B; font-size: 14px;">
            Supported formats: PDF, DOCX, PNG, JPG and JPEG
        </p>
    </div>
    """,
    unsafe_allow_html=True
)

uploaded_file = st.file_uploader(
    "Choose a file from your device",
    type=["pdf", "docx", "png", "jpg", "jpeg"],
    help="Upload a readable medical report to get a simple explanation.",
    key="medical_report_upload"
)

if uploaded_file is not None:
    if st.button("Simplify Medical Report"):
        try:
            # Keep your existing report-processing code here

            extracted_text = []
            file_name = uploaded_file.name.lower()
            file_bytes = uploaded_file.getvalue()

            if file_name.endswith(".docx"):
                document = Document(io.BytesIO(file_bytes))
                extracted_text = [
                    p.text for p in document.paragraphs
                    if p.text.strip()
                ]
                for table in document.tables:
                    for row in table.rows:
                        extracted_text.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

            elif file_name.endswith((".png", ".jpg", ".jpeg")):
                image = Image.open(io.BytesIO(file_bytes))
                extracted_text = [
                    pytesseract.image_to_string(image, lang="eng")
                ]

            elif file_name.endswith(".pdf"):
                with fitz.open(
                    stream=file_bytes, filetype="pdf"
                ) as pdf:
                    for page in pdf:
                        page_text = page.get_text().strip()

                        if not page_text:
                            pix = page.get_pixmap(
                                matrix=fitz.Matrix(2, 2),
                                alpha=False
                            )
                            image = Image.open(
                                io.BytesIO(pix.tobytes("png"))
                            )
                            page_text = pytesseract.image_to_string(
                                image, lang="eng"
                            )

                        extracted_text.append(page_text)

            original_text = "\n".join(extracted_text).strip()

            if not original_text:
                st.warning("No text could be read from this file.")
            else:
                st.subheader("Simplified Explanation")
                explanations = simplify_report(original_text)

                for item in explanations:
                    st.markdown(item)

                simplified_text = "\n".join(
                    item.replace("**", "")
                    for item in explanations
                )

                st.subheader("Basic Number Check")
                warnings = check_numbers(
                    original_text, simplified_text
                )

                
                if warnings:
                    st.warning(
                        "Numbers not found in original report: "
                        + ", ".join(warnings)
                    )
                else:
                    st.success("No unsupported numbers detected.")

                st.download_button(
                    "Download Simplified Report",
                    data=simplified_text,
                    file_name="simplified_report.txt",
                    mime="text/plain"
                )

                with st.expander("View Original Report Text"):
                    st.text_area(
                        "Original Report",
                        value=original_text,
                        height=250,
                        disabled=True
                    )
                    
                st.caption(
                    "Educational prototype only. This app does not "
                    "diagnose medical conditions or guarantee accuracy."
                )

        except Exception as error:
            st.error("Could not process the file: " + str(error))
